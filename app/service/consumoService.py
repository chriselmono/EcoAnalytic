from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import extract, func

from config.db import db
from model.alerta import Alerta
from model.consumo import Consumo


class ConsumoService:

    LOCALIDADES_BOGOTA = (
        "Usaquén", "Chapinero", "Santa Fe", "San Cristóbal", "Usme",
        "Tunjuelito", "Bosa", "Kennedy", "Fontibón", "Engativá", "Suba",
        "Barrios Unidos", "Teusaquillo", "Los Mártires", "Antonio Nariño",
        "Puente Aranda", "La Candelaria", "Rafael Uribe", "Ciudad Bolívar",
        "Sumapaz",
    )
    COLUMNAS_EXCEL = (
        "Localidad", "Barrio", "Consumo_kWh", "Periodo_facturado",
        "Estrato", "Personas_vivienda",
    )
    MAX_FILAS_EXCEL = 1000
    MAX_TAMANO_EXCEL = 10 * 1024 * 1024
    @classmethod
    def crear_consumo(cls, data, user_id):
        datos = cls._validar_datos_consumo(data)
        consumo = cls._crear_modelo_consumo(datos, user_id)
        db.session.add(consumo)
        cls._agregar_alerta_si_aplica(consumo)
        try:
            db.session.commit()
        except Exception:
            db.session.rollback()
            raise
        return {"success": True}

    @staticmethod
    def _parsear_periodo(periodo):
        if isinstance(periodo, datetime):
            valor = periodo.date()
        elif isinstance(periodo, date):
            valor = periodo
        else:
            try:
                valor = datetime.strptime(str(periodo).strip(), "%Y-%m").date()
            except (TypeError, ValueError):
                raise ValueError("El periodo debe tener formato AAAA-MM, por ejemplo 2026-10.")
        valor = valor.replace(day=1)
        if valor < date(2000, 1, 1):
            raise ValueError("El periodo facturado debe ser del año 2000 o posterior.")
        if valor > date.today().replace(day=1):
            raise ValueError("El periodo facturado no puede ser futuro.")
        return valor

    @staticmethod
    def _parsear_consumo(valor):
        if valor is None or isinstance(valor, bool):
            raise ValueError("El consumo en kWh es obligatorio.")
        try:
            consumo = Decimal(str(valor).strip().replace(",", "."))
        except (InvalidOperation, ValueError):
            raise ValueError("El consumo debe ser un número válido.")
        if not consumo.is_finite() or consumo < Decimal("30"):
            raise ValueError("El consumo mínimo permitido es 30 kWh.")
        if consumo > Decimal("99999999.99"):
            raise ValueError("El consumo excede el máximo permitido.")
        if consumo != consumo.quantize(Decimal("0.01")):
            raise ValueError("El consumo admite máximo dos decimales.")
        return consumo.quantize(Decimal("0.01"))

    @staticmethod
    def _parsear_entero(valor, nombre, minimo, maximo):
        if valor is None or str(valor).strip() == "" or isinstance(valor, bool):
            raise ValueError(f"{nombre} es obligatorio y debe ser entero.")
        try:
            numero = Decimal(str(valor).strip())
        except (InvalidOperation, ValueError):
            raise ValueError(f"{nombre} debe ser un número entero.")
        if not numero.is_finite() or numero != numero.to_integral_value():
            raise ValueError(f"{nombre} debe ser un número entero.")
        entero = int(numero)
        if entero < minimo or entero > maximo:
            raise ValueError(f"{nombre} debe estar entre {minimo} y {maximo}.")
        return entero

    @classmethod
    def _validar_datos_consumo(cls, datos):
        localidad = str(datos.get("localidad") or "").strip()
        if localidad not in cls.LOCALIDADES_BOGOTA:
            raise ValueError("Selecciona una localidad válida de Bogotá.")
        barrio = str(datos.get("barrio") or "").strip()
        if len(barrio) < 2 or len(barrio) > 255:
            raise ValueError("El barrio debe tener entre 2 y 255 caracteres.")
        return {
            "localidad": localidad,
            "barrio": barrio,
            "kw_consumo": cls._parsear_consumo(datos.get("kw_consumo")),
            "periodo_consumo": cls._parsear_periodo(datos.get("periodo_consumo")),
            "estrato": cls._parsear_entero(datos.get("estrato"), "El estrato", 1, 6),
            "personas_vivienda": cls._parsear_entero(
                datos.get("personas_vivienda"), "La cantidad de personas", 1, 1000
            ),
        }

    @staticmethod
    def _crear_modelo_consumo(datos, user_id):
        return Consumo(
            localidad=datos["localidad"],
            barrio=datos["barrio"],
            kw_consumo=datos["kw_consumo"],
            periodo_consumo=datos["periodo_consumo"],
            estrato=datos["estrato"],
            personas_vivienda=datos["personas_vivienda"],
            user_id=user_id,
        )

    @staticmethod
    def _agregar_alerta_si_aplica(consumo):
        if consumo.kw_consumo > 500:
            consumo.alertas.append(
                Alerta(
                    titulo="Consumo elevado",
                    descripcion=f"Se registró un consumo de {consumo.kw_consumo} kWh",
                    tipo="danger",
                    user_id=consumo.user_id,
                )
            )

    @classmethod
    def generar_plantilla_excel(cls):
        libro = Workbook()
        hoja = libro.active
        hoja.title = "Consumos"
        hoja.freeze_panes = "A2"
        hoja.auto_filter.ref = "A1:F1001"
        hoja.append(cls.COLUMNAS_EXCEL)

        relleno = PatternFill("solid", fgColor="163B5C")
        for celda in hoja[1]:
            celda.fill = relleno
            celda.font = Font(color="FFFFFF", bold=True)
            celda.alignment = Alignment(horizontal="center")

        for columna, ancho in {"A": 24, "B": 28, "C": 18, "D": 22, "E": 12, "F": 22}.items():
            hoja.column_dimensions[columna].width = ancho
        for fila in range(2, cls.MAX_FILAS_EXCEL + 2):
            hoja.cell(fila, 3).number_format = "0.00"
            # El periodo se captura como texto AAAA-MM, igual que el formulario.
            hoja.cell(fila, 4).number_format = "@"

        validacion_localidad = DataValidation(
            type="list",
            formula1='"' + ",".join(cls.LOCALIDADES_BOGOTA) + '"',
            allow_blank=False,
        )
        validacion_localidad.error = "Selecciona una localidad de la lista."
        validacion_localidad.showErrorMessage = True
        hoja.add_data_validation(validacion_localidad)
        validacion_localidad.add(f"A2:A{cls.MAX_FILAS_EXCEL + 1}")

        validacion_consumo = DataValidation(
            type="decimal", operator="between", formula1="30",
            formula2="99999999.99", allow_blank=False,
        )
        validacion_consumo.error = "Ingresa un consumo desde 30 kWh."
        validacion_consumo.showErrorMessage = True
        hoja.add_data_validation(validacion_consumo)
        validacion_consumo.add(f"C2:C{cls.MAX_FILAS_EXCEL + 1}")

        validacion_periodo = DataValidation(
            type="custom",
            formula1="=AND(ISTEXT(D2),LEN(D2)=7,MID(D2,5,1)=\"-\",IFERROR(VALUE(LEFT(D2,4)),0)>=2000,IFERROR(VALUE(LEFT(D2,4)),0)<=YEAR(TODAY()),IFERROR(VALUE(RIGHT(D2,2)),0)>=1,IFERROR(VALUE(RIGHT(D2,2)),0)<=12,IFERROR(DATE(VALUE(LEFT(D2,4)),VALUE(RIGHT(D2,2)),1),DATE(1900,1,1))<=DATE(YEAR(TODAY()),MONTH(TODAY()),1))",
            allow_blank=False,
        )
        validacion_periodo.error = "Ingresa el periodo con formato AAAA-MM y no uses un mes futuro."
        validacion_periodo.prompt = "Ejemplo: 2026-10."
        validacion_periodo.showErrorMessage = True
        validacion_periodo.showInputMessage = True
        hoja.add_data_validation(validacion_periodo)
        validacion_periodo.add(f"D2:D{cls.MAX_FILAS_EXCEL + 1}")

        validacion_estrato = DataValidation(
            type="list", formula1='"1,2,3,4,5,6"', allow_blank=False
        )
        validacion_estrato.error = "Selecciona un estrato entre 1 y 6."
        validacion_estrato.showErrorMessage = True
        hoja.add_data_validation(validacion_estrato)
        validacion_estrato.add(f"E2:E{cls.MAX_FILAS_EXCEL + 1}")

        validacion_personas = DataValidation(
            type="whole", operator="between", formula1="1",
            formula2="1000", allow_blank=False,
        )
        validacion_personas.error = "Ingresa una cantidad entera mayor que cero."
        validacion_personas.showErrorMessage = True
        hoja.add_data_validation(validacion_personas)
        validacion_personas.add(f"F2:F{cls.MAX_FILAS_EXCEL + 1}")

        instrucciones = libro.create_sheet("Instrucciones")
        instrucciones.column_dimensions["A"].width = 105
        instrucciones["A1"] = "Instrucciones para la carga masiva de consumos"
        instrucciones["A1"].font = Font(bold=True, size=14, color="163B5C")
        pasos = [
            "Registra un recibo por fila en la hoja Consumos, desde la fila 2.",
            "No cambies el nombre de la hoja ni los encabezados.",
            "Localidad: selecciona una opción de la lista.",
            "Barrio: escribe entre 2 y 255 caracteres.",
            "Consumo_kWh: mínimo 30, máximo dos decimales.",
            "Periodo_facturado: escribe el periodo como texto AAAA-MM. Ejemplo: 2026-10.",
            "Estrato: selecciona un valor de 1 a 6.",
            "Personas_vivienda: número entero entre 1 y 1000.",
            "Se admiten hasta 1000 filas. Si alguna fila tiene errores, no se importa ninguna.",
            "Los registros se guardan a nombre de la cuenta que inició sesión.",
        ]
        for fila, texto in enumerate(pasos, start=3):
            instrucciones.cell(fila, 1, texto).alignment = Alignment(wrap_text=True)
        salida = BytesIO()
        libro.save(salida)
        salida.seek(0)
        return salida

    @classmethod
    def importar_consumos_excel(cls, archivo, user_id):
        if archivo is None or not getattr(archivo, "filename", ""):
            return {"success": False, "mensaje": "Selecciona un archivo Excel.", "errores": []}
        if not archivo.filename.lower().endswith(".xlsx"):
            return {"success": False, "mensaje": "El archivo debe tener formato .xlsx.", "errores": []}

        try:
            archivo.stream.seek(0, 2)
            tamano = archivo.stream.tell()
            archivo.stream.seek(0)
        except (AttributeError, OSError):
            tamano = 0
        if tamano > cls.MAX_TAMANO_EXCEL:
            return {"success": False, "mensaje": "El archivo no puede superar los 10 MB.", "errores": []}

        try:
            libro = load_workbook(archivo.stream, read_only=True, data_only=True)
        except (BadZipFile, InvalidFileException, OSError, ValueError):
            return {
                "success": False,
                "mensaje": "No se pudo leer el Excel. Descarga la plantilla e inténtalo de nuevo.",
                "errores": [],
            }

        try:
            if "Consumos" not in libro.sheetnames:
                return {
                    "success": False,
                    "mensaje": "El archivo debe incluir la hoja Consumos de la plantilla.",
                    "errores": [],
                }
            hoja = libro["Consumos"]
            fila_encabezados = next(
                hoja.iter_rows(min_row=1, max_row=1, values_only=True), ()
            )
            encabezados = tuple(
                str(valor).strip() if valor is not None else ""
                for valor in fila_encabezados
            )
            if encabezados != cls.COLUMNAS_EXCEL:
                return {
                    "success": False,
                    "mensaje": "Los encabezados no coinciden. Descarga una plantilla nueva y no cambies sus columnas.",
                    "errores": [],
                }

            filas_validas = []
            errores = []
            cantidad_filas = 0
            for numero_fila, fila in enumerate(
                hoja.iter_rows(min_row=2, values_only=True), start=2
            ):
                valores = tuple(fila[:len(cls.COLUMNAS_EXCEL)])
                if not any(valor is not None and str(valor).strip() for valor in valores):
                    continue
                cantidad_filas += 1
                if cantidad_filas > cls.MAX_FILAS_EXCEL:
                    errores.append({
                        "fila": numero_fila,
                        "detalle": "El archivo supera el máximo de 1000 registros.",
                    })
                    break
                valores += (None,) * (len(cls.COLUMNAS_EXCEL) - len(valores))
                datos_excel = dict(zip(cls.COLUMNAS_EXCEL, valores))
                datos = {
                    "localidad": datos_excel["Localidad"],
                    "barrio": datos_excel["Barrio"],
                    "kw_consumo": datos_excel["Consumo_kWh"],
                    "periodo_consumo": datos_excel["Periodo_facturado"],
                    "estrato": datos_excel["Estrato"],
                    "personas_vivienda": datos_excel["Personas_vivienda"],
                }
                try:
                    filas_validas.append(cls._validar_datos_consumo(datos))
                except ValueError as error:
                    errores.append({"fila": numero_fila, "detalle": str(error)})

            if cantidad_filas == 0:
                return {"success": False, "mensaje": "La hoja Consumos no contiene registros.", "errores": []}
            if errores:
                return {
                    "success": False,
                    "mensaje": f"Se encontraron {len(errores)} errores. No se guardó ningún registro.",
                    "errores": errores[:100],
                }

            consumos = [
                cls._crear_modelo_consumo(datos, user_id)
                for datos in filas_validas
            ]
            db.session.add_all(consumos)
            for consumo in consumos:
                cls._agregar_alerta_si_aplica(consumo)
            try:
                db.session.commit()
            except Exception:
                db.session.rollback()
                return {
                    "success": False,
                    "mensaje": "No se pudo guardar la carga en la base de datos. No se importaron registros.",
                    "errores": [],
                }

            return {
                "success": True,
                "mensaje": f"Se importaron {len(consumos)} consumos correctamente.",
                "registros_importados": len(consumos),
            }
        finally:
            libro.close()

    @staticmethod
    def total_kwh(user_id=None):
        consulta = db.session.query(func.sum(Consumo.kw_consumo))
        if user_id is not None:
            consulta = consulta.filter(Consumo.user_id == user_id)
        return float(consulta.scalar() or 0)

    @staticmethod
    def consumo_mensual(user_id=None):
        consulta = db.session.query(
            Consumo.periodo_consumo,
            func.sum(Consumo.kw_consumo),
        ).filter(Consumo.periodo_consumo.isnot(None))
        if user_id is not None:
            consulta = consulta.filter(Consumo.user_id == user_id)

        resultados = (
            consulta.group_by(Consumo.periodo_consumo)
            .order_by(Consumo.periodo_consumo)
            .all()
        )
        return [
            {
                "mes": periodo.strftime("%Y-%m"),
                "kwh": float(kwh or 0),
            }
            for periodo, kwh in resultados
        ]

    @staticmethod
    def consumo_por_zona(user_id=None):
        consulta = db.session.query(
            Consumo.localidad,
            func.sum(Consumo.kw_consumo),
        )
        if user_id is not None:
            consulta = consulta.filter(Consumo.user_id == user_id)

        resultados = consulta.group_by(Consumo.localidad).all()
        zonas = []
        for nombre, consumo in resultados:
            consumo_val = float(consumo or 0)
            zonas.append(
                {
                    "nombre": nombre,
                    "consumo": consumo_val,
                    "capacidad": consumo_val + 200,
                }
            )
        return zonas

    @staticmethod
    def obtener_consumos(user_id):
        return Consumo.query.filter_by(user_id=user_id).all()

    @staticmethod
    def obtener_consumo(id, user_id):
        return Consumo.query.filter_by(id=id, user_id=user_id).first()

    @staticmethod
    def ultimos_consumos(user_id=None):
        consulta = Consumo.query
        if user_id is not None:
            consulta = consulta.filter(Consumo.user_id == user_id)
        return (
            consulta.order_by(Consumo.periodo_consumo.desc())
            .limit(10)
            .all()
        )

    @staticmethod
    def variacion_mensual(user_id=None):
        hoy = datetime.now()
        mes_actual = hoy.month
        año_actual = hoy.year
        if mes_actual == 1:
            mes_pasado = 12
            año_pasado = año_actual - 1
        else:
            mes_pasado = mes_actual - 1
            año_pasado = año_actual

        filtros_actuales = [
            extract("month", Consumo.periodo_consumo) == mes_actual,
            extract("year", Consumo.periodo_consumo) == año_actual,
        ]
        filtros_pasados = [
            extract("month", Consumo.periodo_consumo) == mes_pasado,
            extract("year", Consumo.periodo_consumo) == año_pasado,
        ]
        if user_id is not None:
            filtros_actuales.append(Consumo.user_id == user_id)
            filtros_pasados.append(Consumo.user_id == user_id)

        actual = (
            db.session.query(func.sum(Consumo.kw_consumo))
            .filter(*filtros_actuales)
            .scalar()
            or 0
        )
        pasado = (
            db.session.query(func.sum(Consumo.kw_consumo))
            .filter(*filtros_pasados)
            .scalar()
            or 0
        )

        actual = float(actual)
        pasado = float(pasado)
        if pasado == 0:
            return "0%"

        variacion = ((actual - pasado) / pasado) * 100
        return f"{round(variacion, 1)}%"

    @staticmethod
    def obtener_alertas(user_id=None):
        consulta = Consumo.query.filter(Consumo.kw_consumo > 500)
        if user_id is not None:
            consulta = consulta.filter(Consumo.user_id == user_id)
        consumos = consulta.order_by(Consumo.periodo_consumo.desc()).all()

        return [
            {
                "tipo": "danger",
                "titulo": "Consumo elevado",
                "desc": f"{consumo.kw_consumo} kWh registrados",
                "hora": (
                    consumo.periodo_consumo.strftime("%Y-%m")
                    if consumo.periodo_consumo
                    else "Periodo no registrado"
                ),
                "zona": consumo.localidad,
            }
            for consumo in consumos
        ]

    @staticmethod
    def obtener_mis_consumos(user_id):
        return (
            Consumo.query.filter_by(user_id=user_id)
            .order_by(Consumo.periodo_consumo.desc())
            .all()
        )

    @classmethod
    def eliminar_consumo(cls, id, user_id):
        try:
            consumo = Consumo.query.filter_by(id=id, user_id=user_id).first()
            if not consumo:
                return {
                    "error": "El registro no existe o no tienes permisos para eliminarlo"
                }

            db.session.delete(consumo)
            db.session.commit()
            return {"success": "Consumo eliminado correctamente"}
        except Exception as error:
            db.session.rollback()
            return {"error": f"Error al eliminar en la base de datos: {error}"}
